#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Gera o briefing do Flow de Winback (Reativacao Depositante) em .docx.

Cores da marca Aposta1:
- Roxo Escuro #160F34 (titulo / cabecalho)
- Roxo #371481 (titulos de secao)
- Verde #00DA6D (acentos, sublinhado, caixa de decisoes)
"""

from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# Paleta da marca
ROXO_ESCURO = RGBColor(0x16, 0x0F, 0x34)
ROXO = RGBColor(0x37, 0x14, 0x81)
VERDE = RGBColor(0x00, 0xDA, 0x6D)
BRANCO = RGBColor(0xFF, 0xFF, 0xFF)
CINZA = RGBColor(0x33, 0x33, 0x33)

ROXO_ESCURO_HEX = "160F34"
VERDE_HEX = "00DA6D"


def set_cell_bg(cell, hex_color):
    """Define a cor de fundo de uma celula de tabela."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tc_pr.append(shd)


def set_cell_border(cell, hex_color, size="18", side="left"):
    """Adiciona uma borda colorida a uma celula (usada como acento)."""
    tc_pr = cell._tc.get_or_add_tcPr()
    borders = tc_pr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tc_pr.append(borders)
    edge = OxmlElement("w:{}".format(side))
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), size)
    edge.set(qn("w:space"), "0")
    edge.set(qn("w:color"), hex_color)
    borders.append(edge)


def add_section_heading(doc, text):
    """Titulo de secao em Roxo com sublinhado/acento verde."""
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.bold = True
    run.font.size = Pt(14)
    run.font.color.rgb = ROXO
    run.font.name = "Calibri"
    p_pr = p._p.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "12")
    bottom.set(qn("w:space"), "3")
    bottom.set(qn("w:color"), VERDE_HEX)
    p_bdr.append(bottom)
    p_pr.append(p_bdr)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p


def add_body(doc, text, bold_label=None):
    """Paragrafo de corpo. bold_label vira o rotulo inicial em negrito."""
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(6)
    if bold_label:
        r = p.add_run(bold_label)
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = CINZA
        r.font.name = "Calibri"
    r2 = p.add_run(text)
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = CINZA
    r2.font.name = "Calibri"
    return p


def add_bullet(doc, text, bold_label=None):
    """Item de lista com bullet. bold_label vira o rotulo em negrito (Roxo)."""
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(4)
    if bold_label:
        r = p.add_run(bold_label)
        r.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = ROXO
        r.font.name = "Calibri"
    r2 = p.add_run(text)
    r2.font.size = Pt(10.5)
    r2.font.color.rgb = CINZA
    r2.font.name = "Calibri"
    return p


def build():
    doc = Document()

    for section in doc.sections:
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # --- Bloco de titulo (fundo Roxo Escuro) ---
    title_table = doc.add_table(rows=1, cols=1)
    title_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = title_table.rows[0].cells[0]
    set_cell_bg(cell, ROXO_ESCURO_HEX)
    set_cell_border(cell, VERDE_HEX, size="24", side="bottom")

    tp = cell.paragraphs[0]
    tp.paragraph_format.space_before = Pt(8)
    tp.paragraph_format.space_after = Pt(2)
    tr = tp.add_run("Briefing: Flow de Winback (Reativação Depositante)")
    tr.bold = True
    tr.font.size = Pt(18)
    tr.font.color.rgb = BRANCO
    tr.font.name = "Calibri"

    sp = cell.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(8)
    sr = sp.add_run("Aposta1 · CRM Lifecycle · 18/08")
    sr.font.size = Pt(11)
    sr.font.color.rgb = VERDE
    sr.font.name = "Calibri"

    doc.add_paragraph().paragraph_format.space_after = Pt(2)

    # --- Objetivo ---
    add_section_heading(doc, "Objetivo")
    add_body(doc,
             "Reativar depositantes que pararam de jogar antes que esfriem de vez. "
             "Hoje 97,3% da base está fora da janela de 7 dias e não existe nenhum "
             "winback rodando, então os players saem sem que ninguém fale com eles.")

    # --- Quem entra (e por que) ---
    add_section_heading(doc, "Quem entra (e por quê)")
    add_body(doc,
             "O gatilho é comportamental, não uma lista estática. O player entra no "
             "flow quando cruza X dias sem apostar ou depositar. A proposta é recortar "
             "por valor, porque a mensagem e a oferta mudam conforme o perfil.")
    add_bullet(doc,
               "cerca de 3.393 hoje. Depositaram há pouco, ainda quentes, alto valor. "
               "Maior ROI, é onde a reativação realmente acontece.",
               bold_label="Trilha A, Crítico high-value (8-30d sem atividade): ")
    add_bullet(doc,
               "cerca de 13.660 elegíveis a SMS. Mais frios, o email entrega pior "
               "aqui, então o SMS puxa melhor.",
               bold_label="Trilha B, Alto (31-90d): ")
    add_bullet(doc,
               "conversão baixa, tratamos em wave dedicada depois.",
               bold_label="Frio 90d+ (126k) fica fora deste flow: ")
    add_body(doc,
             "Selecionamos essa galera porque é o ponto de maior retorno: gente que já "
             "provou que deposita, saiu há pouco, e ainda dá pra trazer de volta com "
             "custo baixo.")

    # --- O que oferecer ---
    add_section_heading(doc, "O que oferecer (decisão que preciso de vocês)")
    add_body(doc,
             "Aqui eu não defino o valor de bônus sozinho. Precisamos fechar com você "
             "e a Marília:")
    add_bullet(doc,
               "oferta mais generosa vale a pena (são high-value, o LTV justifica). "
               "Bônus de reativação, freebet, free spins ou cashback no primeiro "
               "depósito de volta?",
               bold_label="Trilha A: ")
    add_bullet(doc,
               "oferta mais enxuta, com foco em CTA rápido via SMS.",
               bold_label="Trilha B: ")
    add_bullet(doc,
               "Precisa passar pelo time de ofertas e compliance (SIGAP) antes de "
               "fechar o valor.")

    # --- Estrutura do flow ---
    add_section_heading(doc, "Estrutura do flow (proposta)")
    add_bullet(doc,
               "Entrada por trigger comportamental, aguarda o gatilho de dias sem "
               "atividade.")
    add_bullet(doc,
               "reconhecimento mais a oferta principal.",
               bold_label="Mensagem 1 (dia 0): ")
    add_bullet(doc,
               "Espera de 2 a 3 dias, checa se reativou. Se depositou ou apostou, sai "
               "do flow.")
    add_bullet(doc,
               "reforço da oferta e urgência.",
               bold_label="Mensagem 2: ")
    add_bullet(doc,
               "última tentativa, canal alternativo.",
               bold_label="Espera, depois Mensagem 3: ")
    add_bullet(doc,
               "Exit automático em qualquer reativação. Quem não converte cai pro "
               "bucket seguinte.")

    # --- Antes de ligar ---
    add_section_heading(doc, "Antes de ligar (pendências técnicas)")
    add_bullet(doc,
               "Corrigir o segment 326 (event id 166 para 7) e os segments zerados, "
               "senão a elegibilidade entra torta.")
    add_bullet(doc,
               "Resolver os duplicados \"Casino\" e \"Cassino\" pra ninguém receber "
               "mensagem dobrada.")
    add_bullet(doc,
               "Definir a janela exata de dias de cada gatilho com a Marília.")

    # --- Caixa de decisoes (destaque verde) ---
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    dec_table = doc.add_table(rows=1, cols=1)
    dec_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    dcell = dec_table.rows[0].cells[0]
    set_cell_bg(dcell, VERDE_HEX)
    set_cell_border(dcell, ROXO_ESCURO_HEX, size="6", side="left")

    dtitle = dcell.paragraphs[0]
    dtitle.paragraph_format.space_before = Pt(6)
    dtitle.paragraph_format.space_after = Pt(4)
    dtr = dtitle.add_run("Decisão da Marília / vocês")
    dtr.bold = True
    dtr.font.size = Pt(13)
    dtr.font.color.rgb = ROXO_ESCURO
    dtr.font.name = "Calibri"

    dintro = dcell.add_paragraph()
    dintro.paragraph_format.space_after = Pt(4)
    dir_ = dintro.add_run("Dois pontos que só vocês decidem, e que travam o resto:")
    dir_.font.size = Pt(10.5)
    dir_.font.color.rgb = ROXO_ESCURO
    dir_.font.name = "Calibri"

    for txt in [
        "1. Qual a oferta de cada trilha (A e B).",
        "2. A janela exata de dias do gatilho de cada trilha.",
    ]:
        dp = dcell.add_paragraph()
        dp.paragraph_format.space_after = Pt(2)
        dp.paragraph_format.left_indent = Inches(0.15)
        dpr = dp.add_run(txt)
        dpr.bold = True
        dpr.font.size = Pt(10.5)
        dpr.font.color.rgb = ROXO_ESCURO
        dpr.font.name = "Calibri"

    dclose = dcell.add_paragraph()
    dclose.paragraph_format.space_before = Pt(2)
    dclose.paragraph_format.space_after = Pt(6)
    dcr = dclose.add_run("Fechado isso, o flow é montagem nossa.")
    dcr.italic = True
    dcr.font.size = Pt(10.5)
    dcr.font.color.rgb = ROXO_ESCURO
    dcr.font.name = "Calibri"

    doc.save("briefing-winback.docx")
    print("Gerado: briefing-winback.docx")


if __name__ == "__main__":
    build()
